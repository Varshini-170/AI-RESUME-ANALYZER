"""
resume_parser.py
Module 1: Resume Upload (validation helpers)
Module 2: Text Extraction
- Extract text from every PDF page or DOCX paragraph
"""

import os
from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_MB = 5


def validate_file(filepath: str) -> tuple[bool, str]:
    """Validate file type and size before processing."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        return False, f"Unsupported file type '{ext}'. Please upload a PDF or DOCX file."

    size_mb = os.path.getsize(filepath) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        return False, f"File too large ({size_mb:.1f} MB). Max allowed is {MAX_FILE_SIZE_MB} MB."

    return True, "OK"


def extract_text_from_pdf(filepath: str) -> str:
    """Extract text from every page of a PDF resume."""
    reader = PdfReader(filepath)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(filepath: str) -> str:
    """Extract text from every paragraph of a DOCX resume."""
    doc = Document(filepath)
    text_parts = [para.text for para in doc.paragraphs if para.text.strip()]

    # Also capture text inside tables (common in resumes)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_parts.append(cell.text)

    return "\n".join(text_parts)


def extract_resume_text(filepath: str) -> str:
    """Dispatch extraction based on file extension."""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(filepath)
    elif ext == ".docx":
        return extract_text_from_docx(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        path = sys.argv[1]
        ok, msg = validate_file(path)
        print(msg)
        if ok:
            print(extract_resume_text(path)[:500])
